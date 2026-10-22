"""导出：Excel 课表、JSON、ICS 日历。"""
from __future__ import annotations

import json
import uuid
import zlib
from collections import defaultdict
from datetime import date, datetime, time as dtime, timedelta, timezone

from .models import (
    Course,
    DAY_NAMES,
    DEFAULT_PERIOD_TIMES,
    DEFAULT_TERM_WEEKS,
)

_PALETTE = [
    "FFF3E0", "E8F5E9", "E3F2FD", "FCE4EC",
    "F3E5F5", "E0F7FA", "FFFDE7", "EFEBE9",
]


def _periods_text(course: Course) -> str:
    if course.period_start == course.period_end:
        return f"{course.period_start}节"
    return f"{course.period_start}-{course.period_end}节"


def _course_block_text(course: Course) -> str:
    lines = [course.name]
    if course.teacher:
        lines.append(course.teacher)
    if course.location:
        lines.append(course.location)
    lines.append(course.weeks_label)
    return "\n".join(lines)


def _course_fill(course: Course) -> str:
    return _PALETTE[zlib.crc32(course.name.encode("utf-8")) % len(_PALETTE)]


def export_xlsx(courses: list[Course], path: str, period_times=None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    period_times = period_times or DEFAULT_PERIOD_TIMES
    max_period = max([c.period_end for c in courses] + [12])
    max_period = max(8, min(max_period, 20))

    wb = Workbook()
    ws = wb.active
    ws.title = "课表"

    thin = Side(style="thin", color="B0BEC5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="37474F")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    block_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = ["节次", "时间"] + DAY_NAMES
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[1].height = 24
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 13
    for col in range(3, 10):
        ws.column_dimensions[get_column_letter(col)].width = 20

    for p in range(1, max_period + 1):
        row = p + 1
        ws.row_dimensions[row].height = 38
        cell = ws.cell(row=row, column=1, value=p)
        cell.alignment = center
        cell.border = border
        start, end = period_times.get(p, ("", ""))
        cell = ws.cell(row=row, column=2, value=f"{start}-{end}" if start else "")
        cell.alignment = center
        cell.border = border
        cell.font = Font(size=9, color="607D8B")
        for col in range(3, 10):
            ws.cell(row=row, column=col).border = border

    slots: dict[tuple[int, int], list[Course]] = defaultdict(list)
    for course in courses:
        for p in range(course.period_start, min(course.period_end, 20) + 1):
            if 1 <= course.day <= 7:
                slots[(course.day, p)].append(course)

    covered: set[tuple[int, int]] = set()

    def apply_block_style(r1: int, c1: int, r2: int, c2: int, course: Course) -> None:
        fill = PatternFill("solid", fgColor=_course_fill(course))
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                cell = ws.cell(row=rr, column=cc)
                cell.fill = fill
                cell.border = border
                cell.alignment = block_align
                cell.font = Font(size=9, color="263238")

    for day in range(1, 8):
        for p in range(1, max_period + 1):
            key = (day, p)
            if key in covered:
                continue
            slot_courses = slots.get(key)
            if not slot_courses:
                continue
            unique: list[Course] = []
            for course in slot_courses:
                if course not in unique:
                    unique.append(course)

            row = p + 1
            col = day + 2
            course = unique[0] if len(unique) == 1 else None

            if course and course.period_start == p:
                end_period = min(course.period_end, max_period)
                can_merge = True
                for pp in range(p, end_period + 1):
                    other = (day, pp)
                    others = slots.get(other, [])
                    same_course = bool(others) and all(
                        oc.key() == course.key() for oc in others
                    )
                    if other in covered or not same_course:
                        can_merge = False
                        break
                if can_merge and end_period > p:
                    ws.merge_cells(
                        start_row=row,
                        start_column=col,
                        end_row=row + end_period - p,
                        end_column=col,
                    )
                    apply_block_style(row, col, row + end_period - p, col, course)
                    for pp in range(p, end_period + 1):
                        covered.add((day, pp))
                else:
                    if can_merge:
                        covered.add(key)
                    else:
                        covered.add(key)
                    apply_block_style(row, col, row, col, course)
                ws.cell(row=row, column=col, value=_course_block_text(course))
            else:
                text = "\n\n".join(_course_block_text(c) for c in unique)
                starts = {(c.period_start, c.period_end) for c in unique}
                if len(starts) == 1:
                    start_p, end_p = next(iter(starts))
                    end_p = min(end_p, max_period)
                    group_keys = {c.key() for c in unique}
                    same_group = start_p == p and end_p > p and all(
                        {oc.key() for oc in slots.get((day, pp), [])} == group_keys
                        for pp in range(p, end_p + 1)
                    )
                    if same_group:
                        ws.merge_cells(
                            start_row=row,
                            start_column=col,
                            end_row=row + end_p - p,
                            end_column=col,
                        )
                        cell = ws.cell(row=row, column=col, value=text)
                        cell.alignment = block_align
                        cell.border = border
                        cell.font = Font(size=9, color="263238")
                        for pp in range(p, end_p + 1):
                            covered.add((day, pp))
                        continue
                covered.add(key)
                cell = ws.cell(row=row, column=col, value=text)
                cell.alignment = block_align
                cell.border = border
                cell.font = Font(size=9, color="263238")

    ws.freeze_panes = "C2"

    # 第二张表：结构化列表
    ws2 = wb.create_sheet("课程列表")
    list_headers = [
        "课程名称", "星期", "节次", "周次", "教师", "教室",
        "开始时间", "结束时间", "备注",
    ]
    for col, title in enumerate(list_headers, 1):
        cell = ws2.cell(row=1, column=col, value=title)
        cell.fill = PatternFill("solid", fgColor="E3F2FD")
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = border
        ws2.column_dimensions[get_column_letter(col)].width = 16
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["F"].width = 18
    ws2.freeze_panes = "A2"

    sorted_courses = sorted(courses, key=lambda c: (c.day, c.period_start, c.name))
    for ri, course in enumerate(sorted_courses, 2):
        start, end = period_times.get(course.period_start, ("", ""))
        values = [
            course.name,
            course.day_name,
            _periods_text(course),
            course.weeks_label,
            course.teacher,
            course.location,
            course.time_start or start,
            course.time_end or end,
            course.note,
        ]
        for ci, value in enumerate(values, 1):
            cell = ws2.cell(row=ri, column=ci, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws2.auto_filter.ref = f"A1:I{max(1, len(sorted_courses) + 1)}"

    wb.save(path)


def export_json(courses: list[Course], path: str, sources=None) -> None:
    data = {
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "count": len(courses),
        "sources": sources or [],
        "courses": [c.to_dict() for c in courses],
    }
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)


def _ics_escape(value: str) -> str:
    return (
        str(value)
        .replace("\\", "\\\\")
        .replace(";", "\\;")
        .replace(",", "\\,")
        .replace("\r", "")
        .replace("\n", "\\n")
    )


def _fold_line(line: str, limit: int = 73) -> str:
    if len(line.encode("utf-8")) <= limit:
        return line
    parts = []
    current = ""
    for ch in line:
        if len((current + ch).encode("utf-8")) > limit:
            parts.append(current)
            current = " " + ch
        else:
            current += ch
    parts.append(current)
    return "\r\n".join(parts)


def export_ics(
    courses: list[Course],
    path: str,
    first_monday: date | str,
    period_times=None,
) -> tuple[date, int]:
    """按学期第一周周一生成日历事件，返回 (对齐后的周一, 事件数)。"""
    period_times = period_times or DEFAULT_PERIOD_TIMES
    if isinstance(first_monday, str):
        first_monday = date.fromisoformat(first_monday.strip())
    first_monday = first_monday - timedelta(days=first_monday.weekday())

    now_stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//TimetableOCR//CN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "X-WR-CALNAME:课程表",
    ]

    event_count = 0
    for course in courses:
        weeks = course.weeks or DEFAULT_TERM_WEEKS
        start_text = course.time_start or period_times.get(
            course.period_start, ("08:00", "08:45")
        )[0]
        end_text = course.time_end or period_times.get(
            course.period_end, period_times.get(course.period_start, ("08:00", "08:45"))
        )[1]
        try:
            start_clock = dtime.fromisoformat(start_text)
            end_clock = dtime.fromisoformat(end_text)
        except ValueError:
            start_clock = dtime(8, 0)
            end_clock = dtime(8, 45)
        start_dt = datetime.combine(date(2000, 1, 1), start_clock)
        end_dt = datetime.combine(date(2000, 1, 1), end_clock)
        if end_dt <= start_dt:
            end_dt = start_dt + timedelta(minutes=45)

        description = f"节次：{_periods_text(course)}；周次：{course.weeks_label}"
        if course.teacher:
            description += f"；教师：{course.teacher}"
        if not course.weeks:
            description += "（周次未识别，按 1-16 周导出）"

        for week in weeks:
            class_date = first_monday + timedelta(
                days=(week - 1) * 7 + (course.day - 1)
            )
            dtstart = datetime.combine(class_date, start_clock)
            dtend = datetime.combine(class_date, end_clock)
            if dtend <= dtstart:
                dtend = dtstart + timedelta(minutes=45)
            lines.extend(
                [
                    "BEGIN:VEVENT",
                    f"UID:{uuid.uuid4()}@timetable-ocr",
                    f"DTSTAMP:{now_stamp}",
                    f"DTSTART:{dtstart.strftime('%Y%m%dT%H%M%S')}",
                    f"DTEND:{dtend.strftime('%Y%m%dT%H%M%S')}",
                    f"SUMMARY:{_ics_escape(course.name)}",
                    f"LOCATION:{_ics_escape(course.location)}",
                    f"DESCRIPTION:{_ics_escape(description)}",
                    "END:VEVENT",
                ]
            )
            event_count += 1

    lines.append("END:VCALENDAR")
    content = "\r\n".join(_fold_line(line) for line in lines) + "\r\n"
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(content)
    return first_monday, event_count
