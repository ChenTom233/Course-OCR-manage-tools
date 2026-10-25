"""Excel 课表解析：支持 .xlsx/.xlsm/.xls，自动区分网格课表和列表课表。"""
from __future__ import annotations

from .models import compress_weeks
from .parsing import (
    GridCell,
    LogFn,
    _dedupe_courses,
    _noop,
    detect_list_header,
    normalize_text,
    parse_grid,
    parse_list_table,
)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value)


def _spans_to_cells(matrix: list[list[str]], spans) -> list[GridCell]:
    """spans: [(r1, c1, rowspan, colspan), ...]（0 起始）。"""
    height = len(matrix)
    width = max((len(r) for r in matrix), default=0)
    covered = [[False] * width for _ in range(height)]
    cells: list[GridCell] = []
    for r1, c1, rowspan, colspan in spans:
        if r1 >= height or c1 >= width:
            continue
        text = matrix[r1][c1] if c1 < len(matrix[r1]) else ""
        cells.append(GridCell(r1, c1, max(1, rowspan), max(1, colspan), text))
        for r in range(r1, min(r1 + rowspan, height)):
            for c in range(c1, min(c1 + colspan, width)):
                if c < len(covered[r]):
                    covered[r][c] = True
    for r in range(height):
        for c in range(len(matrix[r])):
            if normalize_text(matrix[r][c]) and not covered[r][c]:
                cells.append(GridCell(r, c, 1, 1, matrix[r][c]))
    return cells


def _read_xlsx(path: str, log: LogFn) -> list[tuple[str, list[list[str]], list[GridCell]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue
        matrix = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        if not any(any(v for v in row) for row in matrix):
            continue
        spans = [
            (
                rng.min_row - 1,
                rng.min_col - 1,
                rng.max_row - rng.min_row + 1,
                rng.max_col - rng.min_col + 1,
            )
            for rng in ws.merged_cells.ranges
        ]
        cells = _spans_to_cells(matrix, spans)
        sheets.append((ws.title, matrix, cells))
    wb.close()
    return sheets


def _read_xls(path: str, log: LogFn) -> list[tuple[str, list[list[str]], list[GridCell]]]:
    import xlrd

    book = xlrd.open_workbook(path)
    sheets = []
    for sheet in book.sheets():
        if getattr(sheet, "visibility", 0) != 0:
            continue
        matrix = [
            [_cell_text(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        if not any(any(v for v in row) for row in matrix):
            continue
        spans = []
        for rlo, rhi, clo, chi in sheet.merged_cells:
            spans.append((rlo, clo, rhi - rlo, chi - clo))
        cells = _spans_to_cells(matrix, spans)
        sheets.append((sheet.name, matrix, cells))
    return sheets


def _merge_week_occurrences(courses: list) -> list:
    """列表课表常按“每周一行”记录，把同一门课的多个周次合并成一条。"""
    merged: list = []
    by_key: dict[tuple, object] = {}
    for course in courses:
        if not course.weeks:
            merged.append(course)
            continue
        key = (
            course.name,
            course.day,
            course.period_start,
            course.period_end,
            course.teacher,
            course.location,
            course.time_start,
            course.time_end,
            course.note,
        )
        existing = by_key.get(key)
        if existing is None:
            by_key[key] = course
            merged.append(course)
            continue
        existing.weeks = sorted(set(existing.weeks) | set(course.weeks))
        existing.weeks_text = compress_weeks(existing.weeks)
    return merged


def parse_excel(path: str, log: LogFn = _noop) -> list:
    """解析 Excel 课表文件，返回 Course 列表。"""
    lower = path.lower()
    if lower.endswith(".xls"):
        try:
            sheets = _read_xls(path, log)
        except ImportError:
            log("解析 .xls 需要 xlrd，请先执行：pip install xlrd")
            return []
    elif lower.endswith((".xlsx", ".xlsm")):
        sheets = _read_xlsx(path, log)
    else:
        log(f"不支持的 Excel 格式：{path}")
        return []

    if not sheets:
        log("工作簿里没有可读的工作表")
        return []

    results: dict[int, list] = {}
    profiles: dict[int, dict[str, int]] = {}
    pending: list[int] = []

    # 第一遍：识别带表头的列表式工作表，记录其列映射
    for idx, (title, matrix, cells) in enumerate(sheets):
        log(f"工作表「{title}」共 {len(matrix)} 行")
        courses = parse_list_table(matrix, log)
        if courses:
            log(f"  按列表式解析，识别到 {len(courses)} 条课程")
            results[idx] = courses
            header_row, field_cols, _ = detect_list_header(matrix)
            width = max((len(row) for row in matrix), default=0)
            if header_row >= 0 and field_cols:
                profiles.setdefault(width, field_cols)
        else:
            pending.append(idx)

    # 第二遍：无表头的工作表若与已识别表同宽，套用同一列映射
    for idx in pending:
        title, matrix, cells = sheets[idx]
        width = max((len(row) for row in matrix), default=0)
        courses = []
        field_cols = profiles.get(width)
        if field_cols:
            courses = parse_list_table(
                matrix, log, field_cols=field_cols, header_row=-1
            )
            if courses:
                log(f"  按同列结构识别到 {len(courses)} 条课程")
        if not courses:
            courses = parse_grid(matrix, cells, log)
            log(f"  按网格式解析，识别到 {len(courses)} 条课程")
        results[idx] = courses

    all_courses = []
    for idx in range(len(sheets)):
        all_courses.extend(results.get(idx, []))

    all_courses = _merge_week_occurrences(all_courses)
    before = len(all_courses)
    all_courses = _dedupe_courses(all_courses)
    if before != len(all_courses):
        log(f"去除重复记录 {before - len(all_courses)} 条")
    return all_courses
