"""生成用于自测的合成课表样例。"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.models import DEFAULT_PERIOD_TIMES  # noqa: E402

SAMPLES = Path(__file__).resolve().parent / "samples"
SAMPLES.mkdir(parents=True, exist_ok=True)

FONT = "C:/Windows/Fonts/simhei.ttf"


def make_grid_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "课表"

    ws.merge_cells("A1:I1")
    ws["A1"] = "2025-2026学年第一学期课程表"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["节次", "时间", "周一", "周二", "周三", "周四", "周五", "周六", "周日"]
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _border()

    for p in range(1, 13):
        row = p + 2
        start, end = DEFAULT_PERIOD_TIMES[p]
        ws.cell(row=row, column=1, value=p).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.cell(row=row, column=2, value=f"{start}-{end}").alignment = Alignment(
            horizontal="center", vertical="center"
        )
        for col in range(1, 10):
            ws.cell(row=row, column=col).border = _border()

    def put(day: int, p1: int, p2: int, text: str) -> None:
        col = day + 2
        r1 = p1 + 2
        r2 = p2 + 2
        if r2 > r1:
            ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col)
        cell = ws.cell(row=r1, column=col, value=text)
        cell.alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        cell.border = _border()

    put(1, 1, 2, "高等数学\n张三\n1-16周(单)\n教一101")
    put(4, 3, 4, "大学英语\n李四\n1-16周\n教二202")
    put(2, 5, 6, "数据结构\n王五\n1-16周\n实验楼A305")
    put(7, 1, 2, "体育\n赵六\n1-16周\n体育馆")
    put(3, 7, 8, "线性代数\n钱七\n1-8周\n教三-101")

    for col, width in zip("ABCDEFGHI", [6, 13, 18, 18, 18, 18, 18, 18, 18]):
        ws.column_dimensions[col].width = width

    wb.save(path)


def _border():
    from openpyxl.styles import Border, Side

    thin = Side(style="thin", color="999999")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def make_list_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "课程列表"
    headers = ["课程名称", "教师", "星期", "节次", "周次", "教室"]
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)

    rows = [
        ["操作系统", "陈一", "周一", "1-2节", "1-16", "教一201"],
        ["计算机网络", "林二", "周三", "第3-4节", "1-8周", "教二302"],
        ["数据库原理", "郑三", "周五", "5-6节", "9-16周", "实验楼B101"],
        ["音乐鉴赏", "吴四", "周四", "9-10节", "1-16周(单)", "艺术楼101"],
    ]
    for ri, row in enumerate(rows, 2):
        for ci, value in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=value)
    for col, width in zip("ABCDEF", [18, 10, 10, 12, 14, 16]):
        ws.column_dimensions[col].width = width
    wb.save(path)


def make_xls(path: Path) -> None:
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("课表")
    headers = ["课程名称", "任课教师", "星期", "节次", "周次", "上课地点"]
    for ci, title in enumerate(headers):
        ws.write(0, ci, title)
    rows = [
        ["概率论", "孙八", "周二", "1-2节", "1-16周", "教四401"],
        ["大学物理", "周九", "周四", "5-6节", "1-8周", "实验楼C210"],
    ]
    for ri, row in enumerate(rows, 1):
        for ci, value in enumerate(row):
            ws.write(ri, ci, value)
    wb.save(str(path))


def make_pdf(path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(False)
    pdf.add_page()
    pdf.add_font("SimHei", "", FONT)
    pdf.set_font("SimHei", size=14)
    pdf.cell(0, 12, "2025-2026学年第一学期课程表", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("SimHei", size=9)

    w0, wd = 19, 34
    for title in ["节次", "时间"] + ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]:
        pdf.cell(w0 if title in ("节次", "时间") else wd, 9, title, border=1, align="C")
    pdf.ln()

    course_at = {
        (1, 1): "高等数学 张三 1-16周",
        (4, 3): "大学英语 李四 1-16周",
        (2, 5): "数据结构 王五 1-16周",
        (3, 7): "线性代数 钱七 1-8周",
    }
    pdf.set_font("SimHei", size=6)
    for p in range(1, 9):
        start, end = DEFAULT_PERIOD_TIMES[p]
        pdf.cell(w0, 12, str(p), border=1, align="C")
        pdf.cell(w0, 12, f"{start}-{end}", border=1, align="C")
        for day in range(1, 8):
            pdf.cell(wd, 12, course_at.get((day, p), ""), border=1, align="C")
        pdf.ln()
    pdf.output(str(path))


def make_scanned_pdf(text_pdf: Path, path: Path) -> None:
    from fpdf import FPDF
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(text_pdf))
    try:
        image = doc[0].render(scale=3.0).to_pil().convert("RGB")
    finally:
        doc.close()
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(False)
    pdf.add_page()
    pdf.image(image, x=0, y=0, w=297, h=210)
    pdf.output(str(path))


def make_all() -> dict[str, Path]:
    paths = {
        "grid_xlsx": SAMPLES / "网格课表.xlsx",
        "list_xlsx": SAMPLES / "列表课表.xlsx",
        "list_xls": SAMPLES / "列表课表.xls",
        "pdf": SAMPLES / "课表.pdf",
        "scanned_pdf": SAMPLES / "扫描课表.pdf",
    }
    make_grid_xlsx(paths["grid_xlsx"])
    make_list_xlsx(paths["list_xlsx"])
    make_xls(paths["list_xls"])
    make_pdf(paths["pdf"])
    make_scanned_pdf(paths["pdf"], paths["scanned_pdf"])
    return paths


if __name__ == "__main__":
    for name, path in make_all().items():
        print(f"{name}: {path}")
