"""端到端自测：合成样例 -> 解析 -> 导出 -> GUI 启动。"""
from __future__ import annotations

import json
import subprocess
import sys
import time
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from make_samples import make_all  # noqa: E402

from app.excel_parser import parse_excel  # noqa: E402
from app.exporters import export_ics, export_json, export_xlsx  # noqa: E402
from app.parsing import extract_weeks, parse_course_cell  # noqa: E402
from app.pdf_parser import parse_pdf  # noqa: E402

failures: list[str] = []


def check(condition: bool, name: str, detail: str = "") -> None:
    if condition:
        print(f"  [OK] {name}")
    else:
        print(f"  [FAIL] {name} {detail}")
        failures.append(name)


def by_name(courses, name):
    return next((c for c in courses if c.name == name), None)


def main() -> int:
    print("生成样例…")
    samples = make_all()

    print("\n[1] 文本字段解析")
    weeks, label = extract_weeks("1-16周(单)")
    check(weeks == list(range(1, 16, 2)), "单周解析", f"{weeks}")
    check(label == "1-15周(单)", "单周展示文本", label)
    course = parse_course_cell("周二 第3-4节 高等数学 张三 1-16周 教一101")
    check(course is not None, "整行文本解析")
    if course:
        check(course.name == "高等数学", "整行-课程名", course.name)
        check(course.day == 2 and (course.period_start, course.period_end) == (3, 4),
              "整行-星期节次", f"{course.day} {course.period_start}-{course.period_end}")
        check(course.teacher == "张三", "整行-教师", course.teacher)
        check(course.location == "教一101", "整行-教室", course.location)

    log_lines: list[str] = []

    print("\n[2] Excel 网格课表")
    grid = parse_excel(str(samples["grid_xlsx"]), log_lines.append)
    check(len(grid) == 5, "网格课程数", f"{len(grid)}: {[c.name for c in grid]}")
    c = by_name(grid, "高等数学")
    check(c is not None, "网格-高等数学")
    if c:
        check(c.day == 1 and (c.period_start, c.period_end) == (1, 2), "网格-节次", f"{c.period_start}-{c.period_end}")
        check(c.weeks == list(range(1, 16, 2)), "网格-单周", str(c.weeks))
        check(c.teacher == "张三", "网格-教师", c.teacher)
        check(c.location == "教一101", "网格-教室", c.location)
    c = by_name(grid, "大学英语")
    if c:
        check(c.day == 4 and (c.period_start, c.period_end) == (3, 4), "网格-大学英语", f"{c.day} {c.period_start}-{c.period_end}")
        check(c.location == "教二202", "网格-教二202", c.location)
    c = by_name(grid, "数据结构")
    if c:
        check(c.location == "实验楼A305", "网格-实验楼", c.location)
    c = by_name(grid, "体育")
    check(c is not None and c.location == "体育馆", "网格-体育馆", c.location if c else "missing")
    c = by_name(grid, "线性代数")
    if c:
        check(c.weeks == list(range(1, 9)), "网格-周次 1-8", str(c.weeks))

    print("\n[3] Excel 列表课表")
    listed = parse_excel(str(samples["list_xlsx"]), log_lines.append)
    check(len(listed) == 4, "列表课程数", f"{len(listed)}: {[c.name for c in listed]}")
    c = by_name(listed, "操作系统")
    if c:
        check(c.day == 1 and (c.period_start, c.period_end) == (1, 2), "列表-操作系统")
        check(c.weeks == list(range(1, 17)), "列表-无“周”字周次", str(c.weeks))
        check(c.location == "教一201", "列表-教室", c.location)
    c = by_name(listed, "音乐鉴赏")
    if c:
        check(c.weeks == list(range(1, 16, 2)), "列表-单周", str(c.weeks))

    print("\n[4] 旧版 .xls")
    xls_courses = parse_excel(str(samples["list_xls"]), log_lines.append)
    check(len(xls_courses) == 2, ".xls 课程数", f"{len(xls_courses)}: {[c.name for c in xls_courses]}")
    c = by_name(xls_courses, "概率论")
    if c:
        check(c.day == 2 and c.location == "教四401", ".xls 字段", f"{c.day} {c.location}")

    print("\n[5] PDF 文本层课表")
    pdf_courses = parse_pdf(str(samples["pdf"]), log_lines.append)
    check(len(pdf_courses) >= 4, "PDF 课程数", f"{len(pdf_courses)}: {[c.name for c in pdf_courses]}")
    c = by_name(pdf_courses, "高等数学")
    if c:
        check(c.day == 1, "PDF-星期", str(c.day))
        check(c.weeks == list(range(1, 17)), "PDF-周次", str(c.weeks))
    c = by_name(pdf_courses, "线性代数")
    if c:
        check(c.day == 3, "PDF-线性代数星期", str(c.day))

    print("\n[6] 扫描件 OCR")
    start = time.time()
    ocr_courses = parse_pdf(str(samples["scanned_pdf"]), log_lines.append)
    elapsed = time.time() - start
    print(f"  OCR 用时 {elapsed:.1f}s，识别到：{[c.name for c in ocr_courses]}")
    check(len(ocr_courses) >= 1, "OCR 至少识别到一门课", str(len(ocr_courses)))
    c = by_name(ocr_courses, "高等数学")
    if c:
        check(c.day == 1, "OCR-星期", str(c.day))

    print("\n[7] 导出")
    out_dir = Path(__file__).resolve().parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)
    all_courses = grid + listed + xls_courses + pdf_courses

    xlsx_path = out_dir / "课表.xlsx"
    export_xlsx(all_courses, str(xlsx_path))
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    check(wb.sheetnames == ["课表", "课程列表"], "Excel 工作表", str(wb.sheetnames))
    ws = wb["课程列表"]
    check(ws.max_row == len(all_courses) + 1, "Excel 列表行数", f"{ws.max_row}")
    check(wb["课表"]["A2"].value == 1, "Excel 节次列", str(wb["课表"]["A2"].value))

    json_path = out_dir / "课表.json"
    export_json(all_courses, str(json_path), ["samples"])
    data = json.loads(json_path.read_text(encoding="utf-8"))
    check(data["count"] == len(all_courses), "JSON 条数", str(data["count"]))
    check(any(c["name"] == "高等数学" for c in data["courses"]), "JSON 内容")

    ics_path = out_dir / "课程表.ics"
    monday = date(2026, 9, 7)
    aligned, event_count = export_ics(all_courses, str(ics_path), monday)
    check(aligned == monday, "ICS 周一对齐", str(aligned))
    expected = sum(len(c.weeks) if c.weeks else 16 for c in all_courses)
    check(event_count == expected, "ICS 事件数", f"{event_count} != {expected}")
    ics_text = ics_path.read_bytes().decode("utf-8")
    check(ics_text.startswith("BEGIN:VCALENDAR"), "ICS 头")
    check(ics_text.count("BEGIN:VEVENT") == expected, "ICS VEVENT 数")
    check("SUMMARY:高等数学" in ics_text, "ICS 课程名")
    check("DTSTART:20260907T080000" in ics_text, "ICS 首周时间")
    longest = max(len(line.encode("utf-8")) for line in ics_text.split("\r\n"))
    check(longest <= 75, "ICS 行折叠", f"最长 {longest} 字节")

    print("\n[8] GUI 启动冒烟测试")
    code = (
        "from app.gui import TimetableApp; "
        "app = TimetableApp(); "
        "app.after(600, app.destroy); "
        "app.mainloop()"
    )
    result = subprocess.run(
        [sys.executable, "-c", code],
        cwd=str(ROOT), capture_output=True, text=True, timeout=90,
    )
    check(result.returncode == 0, "GUI 启动", result.stderr[-400:])

    print("\n[9] 真实样例 PDF（一个格子多门课、跨页）")
    real_pdf = ROOT / "陈镓珑(2026-2027-1)课表.pdf"
    if real_pdf.exists():
        real = parse_pdf(str(real_pdf), log_lines.append)
        check(len(real) == 28, "真实样例课程数", f"{len(real)}: {[c.name for c in real]}")

        def find(name, day, pstart):
            return next(
                (c for c in real if c.name == name and c.day == day and c.period_start == pstart),
                None,
            )

        c = find("微处理器原理及应用", 1, 1)
        check(c is not None, "真实-周一微处理器")
        if c:
            check(
                c.weeks == [1, 2, 3, 4, 6, 7, 8, 9, 10, 11, 12, 13, 14, 17, 18],
                "真实-周次",
                str(c.weeks),
            )
            check(c.teacher == "王菲" and c.location == "C-317", "真实-教师教室",
                  f"{c.teacher} {c.location}")
        tue_data = [
            c for c in real
            if c.name == "数据通信技术" and c.day == 2 and c.period_start == 1
        ]
        check(any(c.weeks == [12, 14, 15] for c in tue_data), "真实-双周只作用于本段",
              str([c.weeks for c in tue_data]))
        c = find("数字电路与逻辑设计课程设计", 2, 3)
        check(c is not None, "真实-折行课程名")
        c = find("通信原理", 2, 5)
        check(
            c is not None and c.teacher == "蒙肖雷" and c.location == "C-316",
            "真实-跨页续接字段", f"{c.teacher} {c.location}" if c else "missing",
        )
        c = find("通信原理实验", 2, 5)
        check(
            c is not None and c.weeks == list(range(10, 16)) + [17, 18],
            "真实-续页课程", str(c.weeks) if c else "missing",
        )
        c = find("数据通信技术实验", 5, 3)
        check(c is not None and c.weeks == [6, 9, 11, 12, 13, 14],
              "真实-单周只作用于本段", str(c.weeks) if c else "missing")
        c = find("形势与政策", 1, 7)
        check(c is not None and c.weeks == [11, 12, 13, 14], "真实-第三页课程",
              str(c.weeks) if c else "missing")
        check(
            all(c.name and "未识别" not in c.name for c in real)
            and all(1 <= c.day <= 7 and 1 <= c.period_start <= c.period_end <= 12 for c in real),
            "真实-字段范围",
        )
    else:
        print("  （未找到真实样例，跳过）")

    print("\n[10] 真实样例 Excel（多工作表、按周逐条记录）")
    real_xls = ROOT / "渼陂书院2026-2027-1考研课表安排.xls"
    if real_xls.exists():
        real_courses = parse_excel(str(real_xls), log_lines.append)
        check(len(real_courses) == 39, "真实Excel-合并周次后课程数",
              f"{len(real_courses)}: {[c.name for c in real_courses]}")

        def find_excel(name, day, teacher=None, location=None):
            return next(
                (
                    c for c in real_courses
                    if c.name == name and c.day == day
                    and (teacher is None or c.teacher == teacher)
                    and (location is None or c.location == location)
                ),
                None,
            )

        c = find_excel("高等数学", 1, "张钰婷")
        check(c is not None, "真实Excel-高等数学")
        if c:
            check(
                (c.period_start, c.period_end) == (9, 10)
                and c.weeks == [10, 12, 13, 15, 16],
                "真实Excel-周次合并", f"{c.period_start}-{c.period_end} {c.weeks}",
            )
            check(c.location == "5-301" and c.note == "大三提高1班",
                  "真实Excel-场地教学班", f"{c.location} {c.note}")
        c = find_excel("英语 词汇", 3, "董婉")
        check(c is not None and c.weeks == [9], "真实Excel-课程名合并",
              f"{c.name} {c.weeks}" if c else "missing")
        c = find_excel("英语 翻译", 4, "周世清")
        check(c is not None and c.weeks == [13, 15], "真实Excel-隔周课程",
              str(c.weeks) if c else "missing")

        # Sheet3 无表头，列结构与「教师课表」一致，应套用同一映射
        c = find_excel("数学 线性代数", 1, "杨晓涛", "2-402")
        check(c is not None and c.weeks == [11] and c.note == "大三提高1班",
              "真实Excel-无表头工作表",
              f"{c.weeks} {c.note}" if c else "missing")
        c = find_excel("数学 高数", 1, "张钰婷", "2-402")
        check(c is not None and c.weeks == [10, 14], "真实Excel-无表头周次合并",
              str(c.weeks) if c else "missing")
        c = find_excel("英语 写作", 3, "董婉", "2-402")
        check(c is not None and c.weeks == [11], "真实Excel-无表头课程",
              str(c.weeks) if c else "missing")

        total_occurrences = sum(len(c.weeks) for c in real_courses)
        check(total_occurrences == 52, "真实Excel-周次总数", str(total_occurrences))
        cohorts: dict[str, int] = {}
        for c in real_courses:
            if c.note:
                cohorts[c.note] = cohorts.get(c.note, 0) + 1
        check(
            cohorts == {"大三提高1班": 16, "大三提高2班": 17, "大三提高班": 6},
            "真实Excel-班次分布",
            str(cohorts),
        )
        check(
            all("上午" not in c.name and "城轨概论" not in c.name for c in real_courses),
            "真实Excel-忽略隐藏历史表",
        )
    else:
        print("  （未找到真实 Excel 样例，跳过）")

    print("\n================ 结果 ================")
    if failures:
        print(f"失败 {len(failures)} 项：{failures}")
        print("\n解析日志：")
        print("\n".join(log_lines))
        return 1
    print("全部通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
